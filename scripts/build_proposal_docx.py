# -*- coding: utf-8 -*-
"""Build docs/Supplier_Comparison_Agent_Proposal.docx from docs/PROPOSAL.md.

An editable Word version for teammates. Converts the Markdown subset used in
PROPOSAL.md: headings, paragraphs, bullet and numbered lists, tables, rules,
**bold**, *italic*, `code`, <autolinks> and <br>. Headings use Word's built-in
Heading styles, so the navigation pane and a table of contents work.

Usage:  pip install python-docx && python3 scripts/build_proposal_docx.py
"""
import pathlib
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "PROPOSAL.md"
OUT = ROOT / "docs" / "Supplier_Comparison_Agent_Proposal.docx"

NAVY = RGBColor(0x0B, 0x3D, 0x5C)
FONT = "Calibri"
CODE_FONT = "Consolas"
# **bold**, *italic*, `code`, <http://link>
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|<https?://[^>]+>)")


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text, bold=False):
    """Add Markdown inline text to a paragraph as formatted runs."""
    for i, chunk in enumerate(text.split("<br>")):
        if i:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        for part in INLINE.split(chunk.strip() if i else chunk):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                add_inline(paragraph, part[2:-2], bold=True)
            elif part.startswith("`") and part.endswith("`"):
                run = paragraph.add_run(part[1:-1])
                run.font.name = CODE_FONT
                run.font.size = Pt(9)
                run.bold = bold
            elif part.startswith("<http") and part.endswith(">"):
                add_hyperlink(paragraph, part[1:-1], part[1:-1])
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                run = paragraph.add_run(part[1:-1])
                run.italic = True
                run.bold = bold
            else:
                paragraph.add_run(part).bold = bold


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_rule(doc):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "C8D1DA")):
        bottom.set(qn(k), v)
    border.append(bottom)
    p_pr.append(border)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, cells in enumerate([header] + body):
        for c, text in enumerate(cells):
            cell = table.cell(r, c)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            add_inline(para, text, bold=(r == 0))
            for run in para.runs:
                run.font.size = Pt(9)
            if r == 0:
                shade(cell, "E6EEF4")
    # Repeat the header row on each page.
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.color.rgb = NAVY
        st.font.bold = True


def build():
    doc = Document()
    setup(doc)
    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    after_title = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.fullmatch(r"\|?[\s:|-]+\|?", lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        if stripped == "---":
            add_rule(doc)
        elif stripped.startswith("# "):
            doc.add_paragraph(stripped[2:], style="Title")
            after_title = True
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif re.match(r"^- ", stripped):
            add_inline(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", stripped))
        else:
            # A paragraph runs until a blank line or another block starts.
            text = [stripped]
            while (i + 1 < len(lines) and lines[i + 1].strip()
                   and not re.match(r"^(#|\||- |\d+\. |---$)", lines[i + 1].strip())):
                i += 1
                text.append(lines[i].strip())
            para = doc.add_paragraph()
            add_inline(para, " ".join(t if t.endswith("<br>") else t for t in text).replace("<br> ", "<br>"))
            if after_title:
                for run in para.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                after_title = False
        i += 1
    doc.core_properties.title = "Supplier Comparison Agent: Business Proposal"
    doc.core_properties.author = "Team Show Me Your Token (DAG1YLPM)"
    doc.save(OUT)
    print(f"Wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    sys.exit(build())
